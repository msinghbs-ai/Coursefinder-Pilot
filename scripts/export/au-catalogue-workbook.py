#!/usr/bin/env python3
"""AU catalogue export: the Excel workbook (Platform Admin 2 Nov 2026 11:10 §1).

Reads <out>/manifest.json and <out>/csv/*.csv, writes <out>/au-catalogue-export.xlsx with an Index sheet first (each dataset's sheet,
rows, distinct ids, md5 of its CSV, checks) and then one sheet per dataset that fits Excel's limit (1,048,575 data rows; 32,767 characters
per cell). A dataset over a limit stays CSV only and the Index says why. Values are written as text exactly as in the CSV (null stays an
empty cell); control characters Excel cannot hold are removed from the workbook copy only and counted in the Index: the CSV stays the
exact copy. Prints counts only. Needs openpyxl.
"""
import csv, json, os, re, sys

import openpyxl
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

EXCEL_ROWS = 1_048_575
EXCEL_CELL = 32_767
csv.field_size_limit(1 << 30)


def main(out):
    m = json.load(open(os.path.join(out, 'manifest.json')))
    wb = openpyxl.Workbook(write_only=True)
    index = wb.create_sheet('Index')
    plan = []
    for d in m['datasets']:
        path = os.path.join(out, d['file'])
        reason, longest = None, 0
        if d['rows'] > EXCEL_ROWS:
            reason = f"{d['rows']} rows is over Excel's {EXCEL_ROWS}: CSV only"
        else:
            with open(path, newline='', encoding='utf-8') as f:
                for rec in csv.reader(f):
                    longest = max([longest] + [len(v) for v in rec])
            if longest > EXCEL_CELL:
                reason = f'a value of {longest} characters is over Excel\'s {EXCEL_CELL} per cell: CSV only'
        plan.append((d, path, reason))
    index.append(['AU catalogue export', f"country {m['country']}", f"source commit {m.get('source_commit') or ''}",
                  f"snapshot(s): {len(m['snapshots'])}", f"first snapshot {m['snapshots'][0]['now'] if m['snapshots'] else ''}",
                  f"mode {m['production_impact'].get('mode_used')}", f"checks pass: {m.get('all_checks_pass')}"])
    index.append([])
    index.append(['sheet', 'dataset', 'file', 'rows', 'distinct ids', 'id column', 'count in snapshot', 'csv md5', 'in workbook', 'note'])
    cleaned = {}
    sheets = []
    for d, path, reason in plan:
        if reason:
            continue
        ws = wb.create_sheet(d['sheet'][:31])
        n = 0
        with open(path, newline='', encoding='utf-8') as f:
            r = csv.reader(f)
            for rec in r:
                row = []
                for v in rec:
                    if ILLEGAL_CHARACTERS_RE.search(v):
                        cleaned[d['name']] = cleaned.get(d['name'], 0) + 1
                        v = ILLEGAL_CHARACTERS_RE.sub('', v)
                    row.append(v if v != '' else None)
                ws.append(row)
                n += 1
        sheets.append(d['name'])
        print(f"{d['name']}: {n - 1} rows written")
    for d, path, reason in plan:
        note = reason or ''
        if d['name'] in cleaned:
            note = f"{cleaned[d['name']]} cell(s) had control characters removed in the workbook copy; the CSV is exact"
        if d.get('note'):
            note = (note + ' ' if note else '') + d['note']
        index.append([d['sheet'][:31] if not reason else '', d['name'], d['file'], d['rows'], d.get('distinct_ids'), d.get('id_column'),
                      d.get('count_in_snapshot'), d['md5'], 'no' if reason else 'yes', note])
    index.append([])
    index.append(['check', 'pass', 'detail'])
    for c in m['checks']:
        index.append([c['check'], 'yes' if c['pass'] else 'NO', json.dumps(c['detail'])[:EXCEL_CELL]])
    index.append([])
    index.append(['limitations'])
    for line in m.get('limitations', []):
        index.append([line])
    dest = os.path.join(out, 'au-catalogue-export.xlsx')
    wb.save(dest)
    print(f'workbook: {len(sheets)} dataset sheets plus Index; {sum(1 for _, _, r in plan if r)} dataset(s) CSV only')
    m['workbook'] = {'file': 'au-catalogue-export.xlsx', 'sheets': ['Index'] + sheets, 'csv_only': [{'dataset': d['name'], 'reason': r} for d, _, r in plan if r],
                     'control_characters_removed': cleaned}
    json.dump(m, open(os.path.join(out, 'manifest.json'), 'w'), indent=2)


if __name__ == '__main__':
    main(sys.argv[1])
